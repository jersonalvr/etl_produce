import streamlit as st 
import pandas as pd
from supabase import create_client, Client
import openpyxl
from datetime import datetime, time
from dateutil import parser
import re
import logging
from typing import Optional, List
from st_keyup import st_keyup

def show_carga():
    logging.basicConfig(level=logging.INFO)
    logger = logging.getLogger(__name__)

    st.subheader("REGISTRO DEL DESEMBARQUE DE RECURSOS HIDROBIOLÓGICOS PROCEDENTE DEL ÁMBITO MARÍTIMO")

    @st.cache_resource
    def get_supabase_client() -> Client:
        url = st.secrets["connections"]["supabase"]["SUPABASE_URL"]
        key = st.secrets["connections"]["supabase"]["SUPABASE_KEY"]
        return create_client(url, key)

    supabase = get_supabase_client()

    # Definir 'url' y 'key' para su uso posterior
    url = st.secrets["connections"]["supabase"]["SUPABASE_URL"]
    key = st.secrets["connections"]["supabase"]["SUPABASE_KEY"]

    def es_matricula_formal(matricula):
        if not isinstance(matricula, str):
            return False
        patron = r'^CE-\d{5}-[A-Z]{2}$'
        return bool(re.match(patron, matricula))

    def parsear_fecha(fecha):
        if isinstance(fecha, datetime):
            return fecha.strftime("%Y-%m-%d")
        elif isinstance(fecha, str):
            try:
                fecha_dt = parser.parse(fecha, dayfirst=True, fuzzy=True)
                return fecha_dt.strftime("%Y-%m-%d")
            except parser.ParserError:
                fecha_mod = fecha.replace(" ", "/")
                try:
                    fecha_dt = parser.parse(fecha_mod, dayfirst=True, fuzzy=True)
                    return fecha_dt.strftime("%Y-%m-%d")
                except parser.ParserError:
                    return None
        return None

    def parsear_hora(valor_hora):
        """Convierte diferentes formatos de hora a formato HH:mm"""
        if isinstance(valor_hora, time):
            return valor_hora.strftime("%H:%M")
        elif isinstance(valor_hora, datetime):
            return valor_hora.strftime("%H:%M")
        elif isinstance(valor_hora, str):
            try:
                # Intenta parsear la hora en formato 24 horas
                hora_dt = parser.parse(valor_hora)
                return hora_dt.strftime("%H:%M")
            except:
                try:
                    # Si es un número decimal (ejemplo: 13.30), conviértelo a HH:mm
                    hora_decimal = float(valor_hora)
                    horas = int(hora_decimal)
                    minutos = int((hora_decimal % 1) * 60)
                    return f"{horas:02d}:{minutos:02d}"
                except:
                    return None
        return None

    def get_o_crear_especie(supabase, nombre_comun, nombre_cientifico, tipo):
        try:
            logger.info(f"Buscando especie: {nombre_cientifico}")
            response = supabase.table("especie").select("id").eq("nombre_cientifico", nombre_cientifico).execute()
            if response.data:
                return response.data[0]['id']
            else:
                logger.info(f"Creando nueva especie: {nombre_cientifico}")
                response = supabase.table("especie").insert({
                    "nombre_comun": nombre_comun,
                    "nombre_cientifico": nombre_cientifico,
                    "tipo": tipo,
                    "volumen_total": 0  # Inicializar volumen total
                }).execute()
                return response.data[0]['id']
        except Exception as e:
            logger.error(f"Error al procesar especie: {e}")
            st.error(f"Error al procesar especie: {e}")
            return None

    def actualizar_volumen_especie(supabase, especie_id, volumen_adicional):
        """Actualiza el volumen total de una especie"""
        try:
            # Obtener el volumen actual
            response = supabase.table("especie").select("volumen_total").eq("id", especie_id).execute()
            if response.data:
                volumen_actual = float(response.data[0]['volumen_total'] or 0)
                nuevo_volumen = volumen_actual + float(volumen_adicional or 0)
                
                # Actualizar el volumen total
                supabase.table("especie").update({"volumen_total": nuevo_volumen}).eq("id", especie_id).execute()
                logger.info(f"Volumen actualizado para especie {especie_id}: {nuevo_volumen}")
        except Exception as e:
            logger.error(f"Error al actualizar volumen de especie: {e}")

    def get_o_crear_embarcacion(supabase, matricula, nombre, es_formal):
        try:
            if not matricula:
                return None
                
            logger.info(f"Buscando embarcación: {matricula}")
            response = supabase.table("embarcacion").select("id").eq("matricula", matricula).execute()
            if response.data:
                return response.data[0]['id']
            else:
                logger.info(f"Creando nueva embarcación: {matricula}")
                response = supabase.table("embarcacion").insert({
                    "matricula": matricula,
                    "nombre": nombre,
                    "condicion": "FORMAL" if es_formal else "INFORMAL",
                    "es_formal": es_formal
                }).execute()
                return response.data[0]['id']
        except Exception as e:
            logger.error(f"Error al procesar embarcación: {e}")
            st.error(f"Error al procesar embarcación: {e}")
            return None

    def procesar_y_enviar_archivo(uploaded_file):
        errores = []
        try:
            logger.info(f"Procesando archivo: {uploaded_file.name}")
            wb = openpyxl.load_workbook(uploaded_file, data_only=True)

            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                logger.info(f"Procesando hoja: {sheet_name}")

                # Datos básicos
                departamento = sheet["E6"].value or ""
                provincia = sheet["L6"].value or ""
                distrito = sheet["Q6"].value or ""
                recopilador = sheet["Q8"].value or sheet["Q9"].value or ""
                fecha_cell = sheet["L8"].value
                fecha = parsear_fecha(fecha_cell)

                # Crear recopilación
                recopilacion_data = {
                    "nombre_recopilador": recopilador.strip() if isinstance(recopilador, str) else "",
                    "zona_operacion": f"{departamento}, {provincia}, {distrito}",
                    "fecha_registro": fecha
                }

                response_recopilacion = supabase.table("recopilacion").insert(recopilacion_data).execute()
                recopilacion_id = response_recopilacion.data[0]['id']

                # Procesar cada fila de datos (13 a 57)
                for row in range(13, 58):
                    nombre_comun = sheet[f"C{row}"].value or ""
                    nombre_cientifico = sheet[f"D{row}"].value or ""
                    tipo = sheet[f"X{row}"].value or ""
                    destino = sheet[f"V{row}"].value or ""
                    hora_descarga = parsear_hora(sheet[f"O{row}"].value)  # Parsear hora de descarga

                    if nombre_cientifico:
                        try:
                            cantidad = sheet[f"E{row}"].value
                            unidad_medida = sheet[f"F{row}"].value or ""
                            volumen = sheet[f"G{row}"].value
                            aparejo = sheet[f"H{row}"].value or ""
                            procedencia = sheet[f"I{row}"].value or ""
                            embarcacion = sheet[f"J{row}"].value or ""
                            matricula = sheet[f"K{row}"].value or ""
                            es_formal = es_matricula_formal(matricula)

                            # Obtener o crear especie
                            especie_id = get_o_crear_especie(
                                supabase,
                                nombre_comun.strip() if isinstance(nombre_comun, str) else "",
                                nombre_cientifico.strip() if isinstance(nombre_cientifico, str) else "",
                                tipo.strip() if isinstance(tipo, str) else ""
                            )

                            # Actualizar volumen total de la especie
                            if especie_id and volumen:
                                actualizar_volumen_especie(supabase, especie_id, volumen)

                            # Obtener o crear embarcación
                            embarcacion_id = get_o_crear_embarcacion(
                                supabase,
                                matricula.strip() if isinstance(matricula, str) else "",
                                embarcacion.strip() if isinstance(embarcacion, str) else "",
                                es_formal
                            )

                            # Preparar datos de descarga
                            descarga_data = {
                                "nombre_comun": nombre_comun.strip() if isinstance(nombre_comun, str) else "",
                                "nombre_cientifico": nombre_cientifico.strip() if isinstance(nombre_cientifico, str) else "",
                                "tipo": tipo.strip() if isinstance(tipo, str) else "",
                                "destino": destino.strip() if isinstance(destino, str) else "",
                                "cantidad": cantidad if isinstance(cantidad, (int, float)) else None,
                                "unidad_medida": unidad_medida.strip() if isinstance(unidad_medida, str) else "",
                                "volumen": float(volumen) if isinstance(volumen, (int, float)) else None,
                                "aparejo": aparejo.strip() if isinstance(aparejo, str) else "",
                                "procedencia": procedencia.strip() if isinstance(procedencia, str) else "",
                                "embarcacion_id": embarcacion_id,
                                "matricula": matricula.strip() if isinstance(matricula, str) else "",
                                "tripulantes": sheet[f"L{row}"].value if isinstance(sheet[f"L{row}"].value, int) else None,
                                "dias_de_faena": sheet[f"M{row}"].value if isinstance(sheet[f"M{row}"].value, int) else None,
                                "horas_de_faena": sheet[f"N{row}"].value if isinstance(sheet[f"N{row}"].value, int) else None,
                                "hora_de_descarga": hora_descarga,  # Usar la hora parseada
                                "es_formal": es_formal,
                                "recopilacion_id": recopilacion_id,
                                "registro": f"{recopilacion_id}-{matricula}-{row}"
                            }

                            # Insertar descarga
                            response_descarga = supabase.table("descarga").insert(descarga_data).execute()
                            descarga_id = response_descarga.data[0]['id']

                            # Crear relaciones
                            if especie_id and descarga_id:
                                descarga_especie_data = {
                                    "descarga_id": descarga_id,
                                    "especie_id": especie_id,
                                    "cantidad": cantidad if isinstance(cantidad, (int, float)) else 0
                                }
                                supabase.table("descarga_especies").insert(descarga_especie_data).execute()

                            if embarcacion_id and especie_id:
                                embarcacion_especie_data = {
                                    "embarcacion_id": embarcacion_id,
                                    "especie_id": especie_id,
                                    "cantidad": cantidad if isinstance(cantidad, (int, float)) else 0,
                                    "es_formal": es_formal
                                }
                                supabase.table("embarcacion_especies").insert(embarcacion_especie_data).execute()

                        except Exception as e:
                            logger.error(f"Error al procesar fila {row}: {e}")
                            errores.append(f"Error en fila {row}: {e}")

        except Exception as e:
            logger.error(f"Error al procesar el archivo {uploaded_file.name}: {e}")
            errores.append(f"Error en archivo {uploaded_file.name}: {e}")

        return errores

    # Interface de carga de archivos
    uploaded_files = st.file_uploader("Sube tus reportes Excel", type=["xlsx"], accept_multiple_files=True)

    if uploaded_files:
        if st.button("Enviar a Supabase"):
            progress_bar = st.progress(0)
            status_text = st.empty()
            
            total_files = len(uploaded_files)
            all_errors = []
            
            for i, uploaded_file in enumerate(uploaded_files):
                status_text.text(f"Procesando archivo {i + 1} de {total_files}: {uploaded_file.name}")
                errores = procesar_y_enviar_archivo(uploaded_file)
                all_errors.extend(errores)
                progress_bar.progress((i + 1) / total_files)

            if all_errors:
                st.error("Se encontraron errores durante el proceso:")
                for error in all_errors:
                    st.error(error)
            else:
                st.success("Todos los archivos fueron procesados y enviados correctamente")
                
            status_text.text("Proceso completado")
            progress_bar.progress(1.0)

    def show_supabase_setup_info():
        """Muestra información de configuración para Supabase"""
        
        setup_sql = """
    create or replace function get_tables()
    returns table (table_name text)
    language sql
    as $$
        select table_name::text
        from information_schema.tables
        where table_schema = 'public'
        and table_type = 'BASE TABLE';
    $$;
    """
        
        with st.expander("ℹ️ Configuración de Supabase", expanded=False):
            st.markdown("""
            ### Pasos para configurar Supabase

            1. **Crear función RPC en Supabase:**
                - Ve al Editor SQL de Supabase
                - Copia y ejecuta el siguiente código:
            """)
            
            # Mostrar el SQL con botón de copiado
            st.code(setup_sql, language='sql')
            
            st.markdown("""
            2. **[Verificar credenciales:](https://supabase.com/dashboard/project/_/settings/api)**
                - URL del proyecto: `Settings -> API -> Project URL`
                - API Key: `Settings -> API -> Project API keys -> anon/public`
                
            3. **Permisos necesarios:**
                - La función necesita acceso a `information_schema.tables`
                - El usuario debe tener permisos para ejecutar la función RPC
                
            4. **Solución de problemas:**
                - Asegúrate de que existan tablas en el esquema público
                - Verifica que la base de datos esté activa
                - Confirma que las políticas de seguridad permitan el acceso
            """)

    def get_supabase_tables(url: str, key: str) -> Optional[List[str]]:
        """Obtiene la lista de tablas disponibles en Supabase"""
        try:
            from supabase import create_client, Client
            
            # Crear cliente de Supabase
            supabase: Client = create_client(url, key)
            
            try:
                # Intenta primero usando RPC
                result = supabase.rpc('get_tables').execute()
                
                if hasattr(result, 'data') and result.data:
                    tables = [table['table_name'] for table in result.data]
                    if tables:
                        return sorted(tables)  # Ordenar las tablas alfabéticamente
            except Exception as rpc_error:
                st.warning(f"Método RPC falló: {str(rpc_error)}")
                
                try:
                    # Si RPC falla, intenta con una consulta SQL directa
                    result = supabase.from_('information_schema.tables')\
                        .select('table_name')\
                        .eq('table_schema', 'public')\
                        .eq('table_type', 'BASE TABLE')\
                        .execute()
                    
                    if hasattr(result, 'data') and result.data:
                        return sorted([table['table_name'] for table in result.data])
                except Exception as sql_error:
                    st.warning(f"Consulta SQL directa falló: {str(sql_error)}")
                    
                    # Último intento usando postgREST
                    try:
                        result = supabase.table('tables').select('*').execute()
                        if hasattr(result, 'data') and result.data:
                            return sorted([table['name'] for table in result.data])
                    except Exception as postgrest_error:
                        st.error(f"Todos los métodos de consulta fallaron: {str(postgrest_error)}")
            
            st.warning("No se encontraron tablas en el esquema público")
            # Mostrar ayuda de configuración
            show_supabase_setup_info()
            return None
                        
        except Exception as e:
            st.error(f"Error al conectar con Supabase: {str(e)}")
            st.write("Detalles del error:", str(e))
            # Mostrar ayuda de configuración
            show_supabase_setup_info()
            return None

    def load_supabase_table(url: str, key: str, table_name: str) -> Optional[pd.DataFrame]:
        """Carga una tabla de Supabase como DataFrame"""
        try:
            from supabase import create_client, Client
            
            # Crear cliente de Supabase
            supabase: Client = create_client(url, key)
            
            # Realizar la consulta a la tabla
            response = supabase.table(table_name).select("*").execute()
            
            if hasattr(response, 'data'):
                df = pd.DataFrame(response.data)
                if not df.empty:
                    return df
                else:
                    st.warning(f"La tabla '{table_name}' está vacía")
                    return None
            else:
                st.error("No se pudieron obtener datos de la tabla")
                return None
                
        except Exception as e:
            st.error(f"Error al cargar la tabla de Supabase: {str(e)}")
            st.write("Detalles del error:", str(e))
            return None

    st.markdown("#### Carga desde Supabase")

    # Inicializar variables de estado
    if 'supabase_tables' not in st.session_state:
        st.session_state.supabase_tables = None
    if 'supabase_connected' not in st.session_state:
        st.session_state.supabase_connected = False

    status_container = st.empty()

    col1, col2 = st.columns([1, 4])

    with col1:
        if st.button(
            "Conectar" if not st.session_state.supabase_connected else "Reconectar",
            key="connect_supabase",
            help="Conectar a Supabase y listar tablas disponibles"
        ):
            with st.spinner("Conectando a Supabase..."):
                tables = get_supabase_tables(
                    url,
                    key
                )
                
                if tables:
                    st.session_state.supabase_tables = tables
                    st.session_state.supabase_connected = True
                    status_container.success("✅ Conexión exitosa a Supabase")
                else:
                    st.session_state.supabase_connected = False
                    status_container.error("❌ No se pudieron obtener las tablas. Verifica tus credenciales.")

    if st.session_state.supabase_connected and st.session_state.supabase_tables:
        table_container = st.container()
        
        with table_container:
            selected_table = st.selectbox(
                "Selecciona una tabla:",
                st.session_state.supabase_tables,
                key="supabase_table_selector"
            )
            
            # Inicializar el DataFrame en session_state si no existe
            if 'current_df' not in st.session_state:
                st.session_state.current_df = None
            
            if st.button("Cargar Tabla", key="load_supabase_table"):
                try:
                    with st.spinner("Cargando datos..."):
                        df = load_supabase_table(
                            url,
                            key,
                            selected_table
                        )
                        if df is not None:
                            st.session_state.current_df = df  # Guardar el DataFrame en session_state
                            st.session_state.er_data = df
                            st.success(f"✅ Tabla '{selected_table}' cargada exitosamente")
                
                except Exception as e:
                    st.error(f"❌ Error al cargar la tabla: {str(e)}")
                    st.write("Detalles del error:", str(e))
            
            # Solo mostrar las opciones de búsqueda y la tabla si hay datos cargados
            if st.session_state.current_df is not None:
                df = st.session_state.current_df  # Usar el DataFrame guardado
                
                # Mostrar el DataFrame y las opciones de búsqueda
                st.write("### Datos cargados")
                st.write(f"Total de registros: {len(df)}")
                
                # Crear dos columnas para la búsqueda
                col1, col2 = st.columns([1, 2])
                
                with col1:
                    columns = df.columns.tolist()
                    selected_column = st.selectbox(
                        "Selecciona la columna para buscar",
                        columns,
                        key="search_column"
                    )
                
                with col2:
                    search_term = st_keyup(
                        "Ingresa el término de búsqueda",
                        key="search_term"
                    )
                
                # Filtrar y mostrar datos
                if search_term:
                    try:
                        filtered = df[df[selected_column].astype(str).str.lower().str.contains(search_term.lower(), na=False)]
                        if not filtered.empty:
                            st.write(f"Se encontraron {len(filtered)} registros para '{search_term}' en la columna '{selected_column}'")
                            st.dataframe(filtered)
                            
                            csv = filtered.to_csv(index=False)
                            st.download_button(
                                label="Descargar resultados filtrados",
                                data=csv,
                                file_name=f"busqueda_{selected_column}_{search_term}.csv",
                                mime="text/csv"
                            )
                        else:
                            st.warning(f"No se encontraron coincidencias para '{search_term}' en la columna '{selected_column}'")
                            st.dataframe(df)  # Mostrar tabla completa si no hay resultados
                    except Exception as e:
                        st.error(f"Error al realizar la búsqueda: {str(e)}")
                        st.dataframe(df)  # Mostrar tabla completa en caso de error
                else:
                    # Mostrar tabla completa cuando no hay término de búsqueda
                    st.dataframe(df)
                    
                    # Botón para descargar tabla completa
                    csv_full = df.to_csv(index=False)
                    st.download_button(
                        label="Descargar tabla completa",
                        data=csv_full,
                        file_name=f"{selected_table}_completo.csv",
                        mime="text/csv"
                    )

    return st.session_state.get('er_data', None)
